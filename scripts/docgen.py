# -*- coding: utf-8 -*-
"""通用 docx 讲义生成器：以模板 docx 为基底（保留样式表/页脚/主题），清空正文后注入新内容。

用法:
    from docgen import DocBuilder
    b = DocBuilder('模板.docx')                  # 默认用学而思讲义格式
    b = DocBuilder('模板.docx', profile_path)    # 或用 analyze 后定制的 format_profile.json
    b.title('期中复习专题'); b.section('一、……'); b.points_header()
    b.point_item('1. 要点', ['说明行'])
    b.banner('经典例题'); b.example(['【例1】题干']); b.figure('fig.png', 4.6)
    b.banner('巩固练习'); b.exercise(['1. 练习']); b.save('输出.docx')

默认格式档案 = 学而思三维培优讲义（宋体/Calibri Regular, 五号, 例题区加粗练习区常规,
标题深蓝#1F3A5F, 知识要点绿#2E7D32, 大栏目空心字体居中）。
"""
import json
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.shared import Cm

W_NS = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'

# 内置默认格式档案（学而思三维培优讲义, 2026-09 逆向确认）
DEFAULT_PROFILE = {
    "ascii_font": "Calibri Regular",
    "east_font": "宋体",
    "banner_font": "三极真喵空心简体",
    "title": {"sz": 28, "color": "1F3A5F", "line": ["288", "auto"]},
    "section": {"color": "1F3A5F"},
    "points": {"color": "2E7D32", "label": "🌱 知识要点"},
    "banner": {"line": ["380", "exact"]},
    "body": {"sz": 21, "indent": 283},
}


class FormatProfile:
    def __init__(self, d=None):
        self.d = json.loads(json.dumps(DEFAULT_PROFILE))  # deep copy
        if d:
            self.d.update(d)

    def __getattr__(self, k):
        try:
            return self.d[k]
        except KeyError:
            raise AttributeError(k)


class DocBuilder:
    def __init__(self, template_path, profile=None):
        """profile: dict / json路径 / None(用默认学而思格式)"""
        if isinstance(profile, str):
            profile = json.load(open(profile))
        self.pf = FormatProfile(profile)
        pf = self.pf
        self._rpr_cache = {
            'title': self._rpr(True, pf.title.get('color'), pf.title.get('sz', 28)),
            'section': self._rpr(True, pf.section.get('color')),
            'points': self._rpr(True, pf.points.get('color')),
            'banner': self._rpr(True, None, east=pf.banner_font),
            'bold': self._rpr(True),
            'plain': self._rpr(False),
        }
        self.doc = Document(template_path)
        body = self.doc.element.body
        for child in list(body):
            if child.tag != qn('w:sectPr'):
                body.remove(child)

    def _rpr(self, bold=False, color=None, sz=None, east=None, italic=False):
        pf = self.pf
        sz = sz or pf.body.get('sz', 21)
        east = east or pf.east_font
        x = (f'<w:rPr><w:rFonts w:hint="default" w:ascii="{pf.ascii_font}" '
             f'w:hAnsi="{pf.ascii_font}" w:eastAsia="{east}" w:cs="{pf.ascii_font}"/>')
        x += '<w:b/><w:bCs/>' if bold else '<w:b w:val="0"/>'
        if italic:
            x += '<w:i/>'
        if color:
            x += f'<w:color w:val="{color}"/>'
        x += f'<w:sz w:val="{sz}"/><w:szCs w:val="{sz}"/>'
        if east == pf.banner_font:
            x += '<w:lang w:val="en-US" w:eastAsia="zh-CN"/>'
        return x + '</w:rPr>'

    @staticmethod
    def _esc(t):
        return t.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')

    def make_p(self, runs, align='left', line=None, indent=None, outline=None):
        """runs: [(text, rpr_xml)]"""
        line = line or ('24', 'atLeast')
        pf = self.pf
        ppr = ('<w:pPr><w:keepNext w:val="0"/><w:keepLines w:val="0"/><w:pageBreakBefore w:val="0"/>'
               '<w:kinsoku/><w:wordWrap/><w:overflowPunct/><w:topLinePunct w:val="0"/>'
               '<w:autoSpaceDE/><w:autoSpaceDN/><w:bidi w:val="0"/><w:adjustRightInd w:val="0"/>'
               '<w:snapToGrid w:val="0"/>'
               f'<w:spacing w:before="0" w:after="0" w:line="{line[0]}" w:lineRule="{line[1]}"/>')
        if indent:
            ppr += f'<w:ind w:left="{indent}"/>'
        ppr += f'<w:jc w:val="{align}"/><w:textAlignment w:val="auto"/>'
        if outline is not None:
            ppr += f'<w:outlineLvl w:val="{outline}"/>'
        ppr += (runs[0][1] if runs else self._rpr_cache['plain']) + '</w:pPr>'
        body = ''.join(f'<w:r>{rpr}<w:t xml:space="preserve">{self._esc(t)}</w:t></w:r>'
                       for t, rpr in runs)
        return parse_xml(f'<w:p {W_NS}>{ppr}{body}</w:p>')

    def _append(self, el):
        body = self.doc.element.body
        sect = body.find(qn('w:sectPr'))
        (sect.addprevious(el) if sect is not None else body.append(el))

    # ---------- 结构元素 ----------
    def title(self, text):
        c = self._rpr_cache['title']
        self._append(self.make_p([(text, c)], align='center',
                      line=tuple(self.pf.title.get('line', ['288', 'auto'])), outline=0))

    def section(self, text):
        self._append(self.make_p([(text, self._rpr_cache['section'])]))

    def points_header(self, text=None):
        t = text or self.pf.points.get('label', '🌱 知识要点')
        self._append(self.make_p([(t, self._rpr_cache['points'])]))

    def banner(self, text):
        self._append(self.make_p([(text, self._rpr_cache['banner'])], align='center',
                      line=tuple(self.pf.banner.get('line', ['380', 'exact']))))

    def subsection(self, text):
        self._append(self.make_p([(text, self._rpr_cache['bold'])]))

    def point_item(self, head, bodies):
        """知识条目: 标题行 + 缩进正文行"""
        self._append(self.make_p([(head, self._rpr_cache['bold'])]))
        for b in bodies:
            self._append(self.make_p([(b, self._rpr_cache['bold'])],
                                      indent=self.pf.body.get('indent', 283)))

    def example(self, lines):
        """例题区（模板惯例: 加粗）"""
        for ln in lines:
            self._append(self.make_p([(ln, self._rpr_cache['bold'])]))

    def exercise(self, lines):
        """练习区（模板惯例: 常规字重）"""
        for ln in lines:
            self._append(self.make_p([(ln, self._rpr_cache['plain'])]))

    def blank(self):
        self._append(self.make_p([(' ', self._rpr_cache['plain'])]))

    def figure(self, img_path, width_cm=None):
        """居中图片段。width_cm 缺省按 300dpi 自然尺寸, 限制 2.2~5.5cm"""
        from PIL import Image
        im = Image.open(img_path)
        if width_cm is None:
            width_cm = max(2.2, min(5.5, im.width * 2.54 / 300.0))
        para = self.doc.add_paragraph()
        para.add_run().add_picture(img_path, width=Cm(width_cm))
        new_ppr = parse_xml(
            f'<w:pPr {W_NS}><w:keepNext w:val="0"/><w:spacing w:before="0" w:after="0" '
            f'w:line="24" w:lineRule="atLeast"/><w:jc w:val="center"/>'
            f'{self._rpr(False)}</w:pPr>')
        old = para._p.find(qn('w:pPr'))
        (para._p.replace(old, new_ppr) if old is not None else para._p.insert(0, new_ppr))
        return para

    def save(self, path):
        self.doc.save(path)
